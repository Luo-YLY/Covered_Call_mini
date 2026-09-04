"""Add the original covered-call versus buy-hold anchor to the Ver4 report table.

The edit is intentionally restricted to the cross-ETF early-close table.  It
preserves all surrounding paragraphs, figures, and other tables in the DOCX.
"""

from copy import deepcopy
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, RGBColor


ROOT = Path(__file__).resolve().parents[3]
REPORT_PATH = ROOT / "outputs" / "ver4_summary_report" / "ver4_work_summary_report.docx"
OUTPUT_PATH = REPORT_PATH.with_name("ver4_work_summary_report_含裸持对比.docx")
DELTA_SUMMARY_PATH = ROOT / "outputs" / "ver4_1_delta_buyback" / "ver4_1_delta_buyback_summary.csv"

STATIC_Q100_RULES = {"ATM_Q100", "D10_Q100", "D20_Q100", "D30_Q100", "D40_Q100", "D50_Q100"}
ETF_NAMES = {
    "159915": "创业板ETF（159915）",
    "510050": "上证50ETF（510050）",
    "510300": "沪深300ETF（510300）",
    "510500": "中证500ETF（510500）",
    "588000": "科创50ETF（588000）",
}
HEADER = "原始备兑\n相对裸持（六组平均）"


def copy_cell_format(source, target) -> None:
    source_tc_pr = source._tc.tcPr
    target_tc_pr = target._tc.tcPr
    if target_tc_pr is not None:
        target._tc.remove(target_tc_pr)
    if source_tc_pr is not None:
        target._tc.insert(0, deepcopy(source_tc_pr))

    source_p = source.paragraphs[0]
    target_p = target.paragraphs[0]
    target_p.clear()
    if source_p._p.pPr is not None:
        target_p._p.append(deepcopy(source_p._p.pPr))


def write_like_reference(target, reference, text: str, color: RGBColor | None = None) -> None:
    copy_cell_format(reference, target)
    paragraph = target.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    if reference.paragraphs[0].runs and reference.paragraphs[0].runs[0]._r.rPr is not None:
        run._r.append(deepcopy(reference.paragraphs[0].runs[0]._r.rPr))
    if color is not None:
        run.font.color.rgb = color


def set_table_widths(table) -> None:
    # Keep the extra comparison column inside the existing portrait-page width.
    widths = [Inches(1.27), Inches(0.96), Inches(0.89), Inches(0.89), Inches(0.84), Inches(1.25)]
    table.autofit = False
    for column, width in zip(table.columns, widths):
        for cell in column.cells:
            cell.width = width
            tc_w = cell._tc.tcPr.tcW
            tc_w.w = width.twips
            tc_w.type = "dxa"
    for grid_col, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
        grid_col.w = width


def buyhold_relative_means() -> dict[str, float]:
    summary = pd.read_csv(DELTA_SUMMARY_PATH, dtype={"etf_code": str})
    summary["etf_code"] = summary["etf_code"].str.zfill(6)
    delta80 = summary.loc[summary["delta_buyback_threshold"].eq(0.8)].copy()
    buyhold = delta80.loc[delta80["strategy_name"].eq("BuyHold")].set_index("etf_code")["baseline_strategy_return_sum"]
    static = delta80.loc[delta80["strategy_name"].isin(STATIC_Q100_RULES)].copy()
    static["relative_to_buyhold"] = static.apply(
        lambda row: row["baseline_strategy_return_sum"] - buyhold.loc[row["etf_code"]], axis=1
    )
    return static.groupby("etf_code")["relative_to_buyhold"].mean().to_dict()


def target_table(document: Document):
    for table in document.tables:
        headers = [cell.text.replace("\n", " ") for cell in table.rows[0].cells]
        if any("Delta 0.80" in header for header in headers) and any("TP80" in header for header in headers):
            return table
    raise RuntimeError("Could not find the cross-ETF early-close results table.")


def main() -> None:
    document = Document(REPORT_PATH)
    table = target_table(document)
    means = buyhold_relative_means()

    headers = [cell.text.replace("\n", " ") for cell in table.rows[0].cells]
    try:
        comparison_column = next(index for index, header in enumerate(headers) if "原始备兑" in header and "裸持" in header)
    except StopIteration:
        comparison_column = len(headers)
        table.add_column(Inches(1.25))

    write_like_reference(table.rows[0].cells[comparison_column], table.rows[0].cells[4], HEADER)
    for row in table.rows[1:]:
        etf_name = row.cells[0].text.strip()
        etf_code = next((code for code, name in ETF_NAMES.items() if name == etf_name), None)
        if etf_code is None:
            raise RuntimeError(f"Unknown ETF label in report table: {etf_name}")
        value = means[etf_code]
        text = f"{value * 100:+.1f}pp"
        color = RGBColor(28, 113, 83) if value >= 0 else RGBColor(181, 65, 59)
        write_like_reference(row.cells[comparison_column], row.cells[4], text, color)

    set_table_widths(table)
    document.save(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
