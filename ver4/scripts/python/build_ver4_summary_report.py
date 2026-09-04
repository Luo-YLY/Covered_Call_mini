from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[3]
OUT_DIR = ROOT / "outputs" / "ver4_summary_report"
ASSET_DIR = OUT_DIR / "assets"
OUT_PATH = OUT_DIR / "ver4_work_summary_report.docx"

BLUE = "0B57D0"
NAVY = "183A5A"
INK = "263445"
MUTED = "637285"
PALE_BLUE = "EAF2FB"
PALE_GOLD = "FFF5E6"
GOLD = "B78642"
LINE = "D7E0EA"
GREEN = "167A59"
RED = "B84A4A"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        if edge in kwargs:
            edge_data = kwargs.get(edge)
            tag = "w:{}".format(edge)
            element = borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                borders.append(element)
            for key in ["val", "sz", "space", "color"]:
                if key in edge_data:
                    element.set(qn("w:{}".format(key)), str(edge_data[key]))


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size=10.5, color=INK, bold=False, italic=False):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def add_text(doc, text, size=10.5, color=INK, bold=False, italic=False,
             after=4, before=0, align=None, style=None, line=1.25):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_run_font(r, size=size, color=color, bold=bold, italic=italic)
    return p


def add_rich(doc, parts, after=4, before=0, line=1.25):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    for text, opts in parts:
        r = p.add_run(text)
        set_run_font(r, **opts)
    return p


def add_heading(doc, number, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(f"{number}. {title}")
    set_run_font(r, size=18, color=NAVY, bold=True)
    return p


def add_subheading(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    set_run_font(r, size=11.5, color=BLUE, bold=True)
    return p


def add_bullet(doc, text, color=INK):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.18
    if p.runs:
        set_run_font(p.runs[0], size=10, color=color)
    else:
        set_run_font(p.add_run(text), size=10, color=color)
    return p


def add_note(doc, text, fill=PALE_BLUE, color=INK):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, left={"val": "single", "sz": 14, "color": BLUE},
                    top={"val": "single", "sz": 2, "color": fill},
                    right={"val": "single", "sz": 2, "color": fill},
                    bottom={"val": "single", "sz": 2, "color": fill})
    set_cell_margins(cell, top=90, bottom=90, start=130, end=130)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.18
    r = p.add_run(text)
    set_run_font(r, size=9.4, color=color)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_metric_strip(doc, items):
    table = doc.add_table(rows=1, cols=len(items))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, (label, value, note) in enumerate(items):
        cell = table.cell(0, i)
        set_cell_shading(cell, PALE_BLUE)
        set_cell_border(cell, top={"val": "single", "sz": 4, "color": "B9D4F5"},
                        bottom={"val": "single", "sz": 4, "color": "B9D4F5"},
                        left={"val": "single", "sz": 4, "color": "B9D4F5"},
                        right={"val": "single", "sz": 4, "color": "B9D4F5"})
        set_cell_margins(cell, top=90, bottom=90, start=100, end=100)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p1 = cell.paragraphs[0]
        p1.paragraph_format.space_after = Pt(1)
        set_run_font(p1.add_run(label), size=8.5, color=MUTED, bold=True)
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(1)
        set_run_font(p2.add_run(value), size=13.5, color=NAVY, bold=True)
        p3 = cell.add_paragraph()
        p3.paragraph_format.space_after = Pt(0)
        set_run_font(p3.add_run(note), size=8, color=MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_figure(doc, file_name, caption, width_cm=17.2):
    path = ASSET_DIR / file_name
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    cp = doc.add_paragraph()
    cp.paragraph_format.space_after = Pt(6)
    cp.paragraph_format.line_spacing = 1.1
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(cp.add_run(caption), size=8.5, color=MUTED, italic=True)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, value in enumerate(headers):
        cell = header.cells[idx]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell, top=75, bottom=75, start=80, end=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p.add_run(value), size=8.7, color="FFFFFF", bold=True)
    for ridx, row_values in enumerate(rows):
        row = table.add_row()
        for idx, value in enumerate(row_values):
            cell = row.cells[idx]
            set_cell_margins(cell, top=70, bottom=70, start=80, end=80)
            if ridx % 2 == 0:
                set_cell_shading(cell, "F6F9FC")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx else WD_ALIGN_PARAGRAPH.LEFT
            color = GREEN if isinstance(value, str) and value.startswith("+") else (RED if isinstance(value, str) and value.startswith("-") else INK)
            set_run_font(p.add_run(str(value)), size=8.6, color=color, bold=(idx == 0))
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_page_break(doc):
    doc.add_page_break()


def configure_doc(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.55)
    section.right_margin = Cm(1.55)
    section.top_margin = Cm(1.35)
    section.bottom_margin = Cm(1.35)
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.6)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.2)

    header = section.header
    p = header.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(p.add_run("长城基金 | Ver4 单ETF备兑策略研究"), size=8.5, color=MUTED, bold=True)
    r = p.add_run("    工作总结报告")
    set_run_font(r, size=8.5, color=GOLD)

    footer = section.footer
    p = footer.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(p.add_run("内部研究草案 | 历史回测不构成收益承诺 | 2026-07-17"), size=8, color=MUTED)


def build():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_doc(doc)

    # Cover
    add_text(doc, "VER4.0 | 工作总结报告", size=10.5, color=GOLD, bold=True, after=10, before=20)
    add_text(doc, "单ETF备兑策略：\n从累计回测到逐期现金流诊断", size=25, color=NAVY, bold=True, after=12, line=1.08)
    add_text(doc, "围绕已有 ETF 底仓客户的真实决策问题，重构回测的展示、记账与策略诊断框架。", size=12, color=MUTED, after=22)
    add_metric_strip(doc, [
        ("研究对象", "5 只 ETF", "159915 / 510050 / 510300 / 510500 / 588000"),
        ("基线规则", "DTE30 / Q100", "ATM、D10-D50；100% 覆盖"),
        ("记账口径", "固定名义本金 100", "逐期独立结算，不复利"),
        ("样本区间", "2022-09-30 至 2026-05-27", "以看板数据生成版本为准"),
    ])
    add_text(doc, "本报告的重点", size=12, color=BLUE, bold=True, after=5, before=15)
    add_bullet(doc, "不再只用 CAGR、Sharpe 等宏观指标解释策略，而是逐期展示权利金、行权/平仓扣减、ETF涨跌与完整策略结果。")
    add_bullet(doc, "把“期权腿现金流”和“完整策略收益”拆开，避免把收到的总权利金直接解释为可保留的收益。")
    add_bullet(doc, "把上涨截断、开仓信号和提前买回改造成可定位到单个交易周期的图表与账本。")
    add_note(doc, "结论边界：Ver4 是单ETF、固定名义本金、收盘价代理的历史诊断框架；图中逐期结果可横向比较和简单相加，但不代表复利净值、可成交收益或产品承诺。", fill=PALE_GOLD, color=INK)
    add_text(doc, "报告使用方式：先看单ETF的逐期证据，再判断该标的是否适合作为备兑底仓；择时与提前买回只作为下一阶段的条件性叠加层。", size=10.2, color=INK, after=0, before=15)

    # 1 Framework + cashflow
    add_page_break(doc)
    add_heading(doc, "1", "Ver4 的研究转向：让收益在每一期都可解释")
    add_text(doc, "此前的复利净值、CAGR 与 Sharpe 容易放大少数行情的影响，也无法直接回答客户最关心的问题：本期收了多少权利金、上涨让渡了多少、ETF走弱时策略有没有提供缓冲。Ver4 因此以单个期权周期作为最小分析单元。", size=10.3, after=7)
    add_subheading(doc, "统一的收益拆解")
    add_rich(doc, [
        ("完整策略当期收益 = ", {"size": 10.5, "color": INK, "bold": True}),
        ("ETF 当期收益 + 最终净期权收益", {"size": 10.5, "color": NAVY, "bold": True}),
        ("；其中最终净期权收益 = 总权利金 - 行权/平仓扣减 - 交易成本。", {"size": 10.5, "color": INK}),
    ], after=5)
    add_rich(doc, [
        ("总权利金需再拆为：", {"size": 10.2, "color": INK, "bold": True}),
        ("开仓内在价值 + 开仓时间价值。", {"size": 10.2, "color": NAVY, "bold": True}),
        ("因此若卖出期权在开仓时已略实值，收到的总权利金并不等价于全部可保留的时间价值收入。", {"size": 10.2, "color": INK}),
    ], after=6)
    add_figure(doc, "01_cash_viewport.png", "图1 现金收入页：以 510300 / ATM / Q100 为例，绿色柱为扣摩擦后的权利金净收入，红色柱为行权或平仓扣减；右侧同时保留滚动12期观察。")
    add_note(doc, "阅读要点：44个期权周期的样本中，现金收入和行权/平仓扣减必须并列看。单独强调“累计收到多少权利金”会掩盖强上涨阶段的机会成本。", fill=PALE_BLUE)

    # 2 Ledger
    add_page_break(doc)
    add_heading(doc, "2", "逐期账本：把策略结果还原为可核验的交易记录")
    add_text(doc, "逐期账本支持勾选一个或多个开仓周期，在固定名义本金下汇总 ETF、总权利金、内在价值、时间价值、行权/平仓扣减与最终净期权收益。每一期重新以名义本金 100 开始，避免上一期盈亏改变下一期的权重。", size=10.3, after=6)
    add_figure(doc, "02_cycle_ledger_viewport.png", "图2 逐期账本页：顶部给出当前周期的 ETF 收盘价格、卖出期权和权利金构成；中部把所选周期的完整收益拆开。")
    add_subheading(doc, "账本带来的两个校验")
    add_bullet(doc, "ETF上涨时，完整策略当期收益可以低于总权利金，因为行权/平仓扣减会抵消部分甚至全部总权利金；这不是本金亏损，而是期权腿为保留底仓上涨所付出的成本。")
    add_bullet(doc, "ETF下跌时，期权腿通常为正并形成缓冲，但完整策略仍可能为负；备兑不是下跌保护策略，而是用权利金降低下跌幅度。")

    # 3 upside surrender
    add_page_break(doc)
    add_heading(doc, "3", "上涨让渡：区分温和上涨与强上涨的代价")
    add_text(doc, "上涨让渡页以 ETF 当期涨跌为横轴、备兑相对裸持的贡献为纵轴。点落在零轴下方时，代表本期策略虽然可能仍为正，但低于同周期仅持有 ETF 的结果。", size=10.3, after=6)
    add_figure(doc, "03_upside_surrender_viewport.png", "图3 上涨让渡页：510300 / ATM / Q100 的单周期散点。强上涨期更容易落入相对裸持为负的区域，方形表示发生行权。")
    add_note(doc, "该图提供的是产品定位的边界，而不是负面结论：备兑更适合希望把部分上涨弹性转换为阶段性现金流、且可接受强上涨时落后于裸持的底仓客户。不同虚值程度会改变让渡阈值与权利金厚度，需逐ETF单独诊断。", fill=PALE_GOLD)

    # 4 timing
    add_page_break(doc)
    add_heading(doc, "4", "择时分析：从事后关系到开仓时可得的信号")
    add_text(doc, "看板提供“开仓IV历史分位”和“开仓IV-RV”两个手动开关：满足条件时卖出期权，不满足时仅持有ETF。IV分位只使用更早开仓日的IV，IV-RV应明确区分 IV 取开仓日、RV 取开仓前一日或此前窗口的时间可得性。", size=10.3, after=6)
    add_figure(doc, "05_timing_viewport.png", "图4 择时分析页：ETF、始终卖出与当前择时情景的逐期收益并列显示；当前截图为开关未启用的基准对照。")
    add_subheading(doc, "当前阶段的判断")
    add_bullet(doc, "IV-RV 是候选开仓信号，而非已验证的盈利规则。它可能解释权利金是否更厚，却不能天然预测随后上涨截断或方向性行情。")
    add_bullet(doc, "后续需要按滚动历史窗口、不同阈值与样本外切分检验，并把“未卖出时只持有ETF”的机会成本纳入同一逐期对照。")
    add_note(doc, "事后 VRP、持有期实现波动率等指标可用于解释，不可直接作为开仓日的可用信号；报告中应明确标注“使用未来信息”。", fill=PALE_GOLD)

    # 5 Early close chart/table
    add_page_break(doc)
    add_heading(doc, "5", "提前平仓：Delta 买回与 TP80 的条件性价值")
    add_text(doc, "Ver4 将提前买回放在原始备兑之上作为覆盖层：触发后买回原卖出认购期权，后续不重开，ETF持有至原周期结束。Delta规则要求开仓后首个满足绝对 Delta 阈值的交易日，且剩余到期日不少于5天；以日终价格代理回购，并计入额外5bp滑点和半个假设价差。", size=10.1, after=6)
    add_figure(doc, "04_early_close_viewport.png", "图5 提前平仓页：灰柱为原始备兑期权腿，深色柱为提前买回后期权腿；浅色阴影标记该周期完整策略结果相对基线提升或下降。")
    add_subheading(doc, "跨ETF的方向性结果（六组静态 Q100 规则的平均变化）")
    add_table(doc,
              ["ETF", "Delta 0.80\n完整策略变化", "TP80\n完整策略变化", "Delta行权\n减少比例", "TP80行权\n减少比例"],
              [
                  ["创业板ETF（159915）", "+20.1pp", "+2.5pp", "71.9%", "13.5%"],
                  ["上证50ETF（510050）", "-10.7pp", "-2.0pp", "71.6%", "11.4%"],
                  ["沪深300ETF（510300）", "-7.5pp", "-1.4pp", "72.3%", "7.2%"],
                  ["中证500ETF（510500）", "+5.9pp", "-1.7pp", "77.4%", "8.6%"],
                  ["科创50ETF（588000）", "+45.8pp", "-3.9pp", "67.9%", "8.9%"],
              ], widths=[3.6, 3.0, 2.5, 2.5, 2.5])
    add_text(doc, "注：以上为相对“原始备兑”的固定名义本金逐期加总差额，不是复利收益率；结果显示 Delta 买回能显著减少行权，但对完整策略的贡献高度依赖 ETF 与行情路径。TP80 触发更频繁，行权减少却有限，当前不宜作为强上涨保护规则。", size=8.7, color=MUTED, after=0)

    # 6 Delta attribution
    add_page_break(doc)
    add_heading(doc, "6", "Delta 触发后的路径归因：为什么同一规则有时有效、有时无效")
    add_text(doc, "单看“是否触发Delta”不足以判断提前买回的质量。看板进一步记录买回后ETF期末涨跌、买回后最大继续上涨与策略变化；如果买回后出现持续上涨，保留ETF上行通常可改善相对基线结果；若价格很快回落，提前回购往往仅锁定了额外成本。", size=10.3, after=6)
    add_figure(doc, "06_delta_path_attribution_viewport.png", "图6 Delta触发后的路径归因：每个点是一笔实际触发的买回。绿色为买回后相对原始备兑改善，红色为下降；该分析使用买回后的未来路径，仅作事后解释。")
    add_subheading(doc, "已观察到的模式")
    add_bullet(doc, "全样本下，Delta-only触发的周期平均变化为正，但同时触发Delta与TP80的周期并不稳定；说明“高Delta”并不自动等同于后续继续上涨。")
    add_bullet(doc, "Delta阈值过高会明显变晚：跨ETF、跨规则平均而言，阈值由0.70提高到0.80、0.90后，完整策略改善依次减弱。")
    add_bullet(doc, "下一步的正确方向不是立即固化阈值，而是用触发后路径特征构建第二道确认条件，并对买回时间、剩余到期日和交易摩擦做敏感性分析。")

    # 7 closing
    add_page_break(doc)
    add_heading(doc, "7", "阶段结论与下一步工作")
    add_subheading(doc, "Ver4 已完成的能力")
    add_bullet(doc, "形成以单ETF、单期权周期为中心的固定名义本金账本，可分别查看总权利金、时间价值、行权/平仓扣减、ETF收益和完整策略收益。")
    add_bullet(doc, "形成上涨让渡、市场情景、持有期路径、事后诊断、择时分析与提前平仓的看板闭环，支持从一张总图回到单个周期核验。")
    add_bullet(doc, "验证提前买回存在标的差异和路径依赖：Delta覆盖层不是通用增益，需要与ETF特征、市场阶段和执行成本联合评估。")
    add_subheading(doc, "下一阶段优先级")
    add_table(doc,
              ["优先级", "工作项", "希望回答的客户问题"],
              [
                  ["P0", "逐ETF的权利金、上涨让渡与下跌缓冲分布；按温和/强上涨和下跌阶段分组。", "我的底仓在什么行情下值得卖Call？"],
                  ["P1", "IV分位、IV-RV的严格时间可得性定义；滚动样本外与阈值稳健性检验。", "什么条件下卖出比始终卖出更合理？"],
                  ["P2", "Delta买回的二次确认、DTE边界、买回后ETF路径与成交摩擦敏感性。", "强上涨发生时，是否值得用成本换回上涨弹性？"],
                  ["P3", "补充分红、合约调整、bid/ask与更高频成交数据；重估IV反解异常。", "这些回测结果与真实可执行结果相差多远？"],
              ], widths=[1.5, 8.0, 5.8])
    add_subheading(doc, "可用于客户沟通的定位")
    add_text(doc, "产品不应承诺“每期稳定赚取权利金”，而应清楚呈现取舍：在大多数非强上涨周期，期权腿为ETF底仓提供可观察的现金流与一定缓冲；在强上涨周期，客户以部分上涨弹性换取这份现金流。Ver4的价值在于把这项取舍逐期、逐标的、可核验地展示出来。", size=10.4, after=7)
    add_note(doc, "风险提示：本报告采用历史收盘数据与模型/假设交易摩擦。未包含真实逐笔成交、完整买卖价、全部分红与合约调整影响；所有结果仅供研究讨论，不构成投资建议、产品收益承诺或实盘执行方案。", fill=PALE_GOLD)

    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build()
