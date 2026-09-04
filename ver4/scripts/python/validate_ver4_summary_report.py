from pathlib import Path
import zipfile

from docx import Document


root = Path(__file__).resolve().parents[3]
report = root / "outputs" / "ver4_summary_report" / "ver4_work_summary_report.docx"
document = Document(report)

with zipfile.ZipFile(report) as archive:
    media = [name for name in archive.namelist() if name.startswith("word/media/")]
    document_xml = archive.read("word/document.xml").decode("utf-8")

print(
    {
        "docx_bytes": report.stat().st_size,
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
        "embedded_images": len(media),
        "explicit_page_breaks": document_xml.count('w:type="page"'),
    }
)
for image in media:
    print(image)
